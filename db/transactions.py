import os.path
import time
from collections import Counter
from copy import copy
from datetime import datetime, timedelta

import openpyxl
import pandas as pd
from openpyxl.reader.excel import load_workbook
from openpyxl.utils import get_column_letter
from psycopg2 import DatabaseError

from app import basic_funcs, app_logger
from app.basic_funcs import rgetattr, xls2xlsx_
from db import session, sp
from db.tables import PRODUCT

logger = app_logger.get_logger(__name__)


def import_file_data(product_name, file, up_node_id):
    total_started_at = time.perf_counter()
    logger.info("WorkData import started: file=%s, up_node_id=%s", file, up_node_id)

    def progress_message(imported, total):
        remaining = max(total - imported, 0)
        return f'Импортировано {imported} из {total}, осталось {remaining}'

    def update_import_progress(imported, total, message=None):
        session.update_progress(
            current=imported,
            total=total,
            message=message or progress_message(imported, total),
            detail=progress_message(imported, total),
        )

    def copy_import_file_data(records):
        logger.info("WorkData import COPY started: records=%s", len(records))
        copy_started_at = time.perf_counter()
        copy_rows = [
            (
                record[1],
                record[2],
                record[3],
                record[4],
                record[5],
                record[6],
                record[7],
                0 if record[9] is None else record[9],
                record[10],
            )
            for record in records
        ]
        logger.debug(
            "WorkData import COPY rows prepared: records=%s, elapsed=%.4fs",
            len(copy_rows),
            time.perf_counter() - copy_started_at,
        )

        total = len(copy_rows)
        update_import_progress(0, total)

        def on_copy_progress(imported, total_records):
            update_import_progress(imported, total_records)

        inserted_count = session.copy_import_file_data(
            copy_rows,
            batch_size=10000,
            progress_callback=on_copy_progress,
        )
        update_import_progress(
            total,
            total,
            message=f'Импорт рабочих данных: импортировано {total} из {total}',
        )
        return inserted_count

    def get_cells_for_import(filename, excel_id, worksheet=None) -> list:
        workbook_started_at = time.perf_counter()
        session.update_loading_bar('Импорт рабочих данных: чтение файла...')
        wb = xls2xlsx_(filename)
        logger.info(
            "WorkData import workbook loaded: elapsed=%.4fs",
            time.perf_counter() - workbook_started_at,
        )
        ws = wb[worksheet] if worksheet else wb.active

        session.update_loading_bar('Импорт рабочих данных: разбор листа...')
        parse_started_at = time.perf_counter()
        # Загружаем данные в pandas DataFrame
        data = pd.DataFrame([[cell.value for cell in row] for row in ws.iter_rows()])
        data = data.dropna(how='all', axis=1)  # Убираем полностью пустые столбцы
        data = data.dropna(how='all', axis=0)  # Убираем полностью пустые строки

        # Проверяем и корректируем индексацию
        row_names = data.iloc[:, 0].dropna().astype(str).values  # Первая колонка — имена строк
        rows = data.iloc[:, 1:]  # Убираем первую колонку, она уже обработана
        non_empty_columns = list(range(rows.shape[1]))  # Все оставшиеся столбцы (от 0 до n-1)
        populated_cells = int(rows.notna().sum().sum())
        logger.info(
            "WorkData import worksheet parsed: rows=%s, columns=%s, populated_cells=%s, elapsed=%.4fs",
            len(row_names),
            len(non_empty_columns),
            populated_cells,
            time.perf_counter() - parse_started_at,
        )

        # Получаем IDs параметров
        session.update_loading_bar('Импорт рабочих данных: подготовка справочника имён...')
        sprav_started_at = time.perf_counter()
        param_ids = sp.add_upd_sprav_names_array(row_names.tolist())
        logger.info(
            "WorkData import sprav names prepared: unique_names=%s, elapsed=%.4fs",
            len(set(row_names.tolist())),
            time.perf_counter() - sprav_started_at,
        )

        # Дата для столбцов с уникальностью
        session.update_loading_bar('Импорт рабочих данных: подготовка записей...')
        records_started_at = time.perf_counter()
        curr_date = datetime.now()
        date_column_dict = {
            i: curr_date + timedelta(milliseconds=i) for i in non_empty_columns
        }  # Уникальное время для каждого столбца

        # Формируем результаты для вставки
        result_columns = [
            (0, excel_id, 0, 'type', date_column_dict[i], None, None, 'column', False, 0, None)
            for i in non_empty_columns
        ]
        result_columns += [
            (0, excel_id, 0, 'column_npp', date_column_dict[i], None, None, f'{i}', False, 0, None)
            for i in non_empty_columns
        ]

        result_rows = [
            (0, excel_id, 0, 'type', None, None, None, 'row', False, 0, param_ids[i])
            for i, name in enumerate(row_names)
        ]
        result_rows += [
            (0, excel_id, 0, 'row_npp', None, None, None, str(i), False, 0, param_ids[i])
            for i, name in enumerate(row_names)
        ]

        # Преобразование всех данных через генераторы
        def process_cells():
            for i, name in enumerate(row_names):
                for j, cell_value in enumerate(rows.iloc[i]):
                    if pd.isna(cell_value):
                        continue
                    yield (
                        0, excel_id, 0, 'value',
                        date_column_dict[j],  # Уникальная дата из date_column_dict
                        None, None, str(cell_value), False, 0, param_ids[i]
                    )
                    yield (
                        0, excel_id, 0, 'accuracy',
                        date_column_dict[j],  # Уникальная дата из date_column_dict
                        None, None, '2', False, 0, param_ids[i]
                    )
                    yield (
                        0, excel_id, 0, 'type',
                        date_column_dict[j],  # Уникальная дата из date_column_dict
                        None, None, 'cell', False, 0, param_ids[i]
                    )

        result_cells = list(process_cells())

        records = result_cells + result_rows + result_columns
        by_prop = Counter(record[3] for record in records)
        by_prop.setdefault('bold', 0)
        by_prop.setdefault('size', 0)
        logger.info(
            "WorkData import records prepared: total_records=%s, by_prop=%s, elapsed=%.4fs",
            len(records),
            dict(sorted(by_prop.items())),
            time.perf_counter() - records_started_at,
        )
        session.update_progress(
            current=0,
            total=len(records),
            message=f'Импорт рабочих данных: подготовлено {len(records)} записей',
            detail=progress_message(0, len(records)),
        )

        return records

    try:
        session.begin_progress('Импорт рабочих данных', blocking=True)
        # TODO нужно ли проверять отдельно каждую функцию?
        product = (None, None, up_node_id, 5, 'name', product_name, None, 0, None)
        new_product = sp.new_update_product_from_record(product)
        version = copy(new_product)
        version.prod_prop = 'final_version'
        version.prod_prop_value = '0'
        version = sp.new_update_product_from_record(version.table_fit(PRODUCT))
        datafile = sp.new_uniskad_datafile(int(new_product.id), 'input_excel', product_name[:30], file, "")
        with open(file, 'rb') as f:
            sp.new_uniskad_binfile(datafile.id_datafile, datafile.full_name, 'xlsx',
                                   '',
                                   '', f.read())
        sp.create_update_import_file_state(datafile.id_datafile, 10, 0)
        records = get_cells_for_import(datafile.full_name, datafile.id_datafile)
        try:
            copy_import_file_data(records)
        except Exception:
            logger.exception("Fast COPY import failed, falling back to new_excel_data_array")
            sp.new_excel_data_array(records)
        session.update_loading_bar('Импорт рабочих данных: обновление таблицы...')
        logger.info(
            "WorkData import completed: total_elapsed=%.4fs",
            time.perf_counter() - total_started_at,
        )
        return new_product, version
    finally:
        session.end_progress()


# @session.transaction
def import_other_file_data(filename, file, file_type, up_node_id):
    # TODO нужно ли проверять отдельно каждую функцию?
    product = (None, None, up_node_id, file_type, 'name', filename, None, 0, None)
    new_product = sp.new_update_product_from_record(product)
    new_product.type_ = 'file'
    datafile = sp.new_uniskad_datafile(int(new_product.id), 'other', filename, file, "")
    with open(file, 'rb') as f:
        sp.new_uniskad_binfile(datafile.id_datafile, datafile.full_name, 'other',
                               '',
                               '', f.read())
    return new_product


# @session.transaction
def import_project_other_file_data(filename, file, file_type, up_node_id):
    # TODO нужно ли проверять отдельно каждую функцию?
    product = (None, None, up_node_id, file_type, 'name', filename, None, None, None, 0, None)
    new_product = sp.new_update_project_from_record(product)
    new_product.type_ = 'file'
    datafile = sp.new_uniskad_datafile_project(int(new_product.id), 'other', filename, file, "")
    with open(file, 'rb') as f:
        sp.new_uniskad_binfile(datafile.id_datafile, datafile.full_name, 'pdf',
                               '',
                               '', f.read())
    return new_product


# @session.transaction
def delete_file_data(product_id, file_type, other_format=False):
    datafile = sp.get_product_uniskad_files(product_id, file_type)
    if not other_format:
        sp.remove_import_file_data(datafile.id_datafile)
        sp.remove_import_file_state(datafile.id_datafile)
    sp.remove_uniskad_datafiles_binary(datafile.id_datafile)
    sp.remove_uniskad_datafiles(datafile.id_datafile)
    sp.delete_product(product_id, True, True, True)
    return True


# @session.transaction
def delete_file_data_project(product_id, file_type, other_format=False):
    datafile = sp.get_project_uniskad_files(product_id, file_type)
    sp.remove_uniskad_datafiles_binary(datafile.id_datafile)
    sp.remove_uniskad_datafiles(datafile.id_datafile)
    sp.delete_project(product_id, True, True, True)
    return True
