import datetime

import pytest

pytest.importorskip('PySide2')

from openpyxl import load_workbook
from PySide2.QtCore import Qt
from PySide2.QtWidgets import QApplication

from app.plugins.work_data.widgets.table import WorkDataTableView
from db.schemas import ImportFileData


@pytest.fixture(scope='module')
def application():
    return QApplication.instance() or QApplication([])


def record(row, column, prop, value, record_id=1):
    return ImportFileData({
        'id_record': record_id,
        'is_secret': False,
        'id_excel_file': 1,
        'file_version': 1,
        'excel_param_name': row,
        'param_prop_name': prop,
        'date_time_izm': column,
        'zamer_n': None,
        'rejim_zamer': None,
        'prop_value': value,
        'deleted': False,
        'id_name': 1,
        'sprav_name': None,
        'accuracy': 2,
        'id_eizm': None,
        'eizm_short': None,
        'eizm_full': None,
    })


def sample_records():
    column_1 = datetime.datetime(2024, 1, 1)
    column_2 = datetime.datetime(2024, 1, 2)
    return [
        record('B', None, 'row_npp', '2'),
        record('B', None, 'accuracy', '1'),
        record('A', None, 'row_npp', '1'),
        record('A', None, 'accuracy', '2'),
        record(None, column_2, 'column_npp', '2'),
        record(None, column_1, 'column_npp', '1'),
        record('A', column_1, 'value', '1.25'),
        record('A', column_1, 'formula', ''),
        record('A', column_2, 'value', '2'),
        record('B', column_1, 'value', '3'),
        record('B', column_2, 'value', '4'),
    ]


def make_view(application):
    view = WorkDataTableView()
    view.model().load_data(sample_records())
    return view


def test_model_lookup_and_compatibility_indexes_are_ordered_and_lazy(application):
    view = make_view(application)
    model = view.model()

    assert model.ord_rows == ['A', 'B']
    assert model.ord_columns == [datetime.datetime(2024, 1, 1), datetime.datetime(2024, 1, 2)]
    assert model._items == {}

    index = model.index(0, 0)
    item = view.itemFromIndex(index)
    assert item.key == ('A', datetime.datetime(2024, 1, 1))
    assert view.indexFromItem(item) == index
    assert len(model._items) == 1
    assert model.data(index, Qt.DisplayRole) == '1.25'


def test_set_data_updates_cell_and_save_queue(application):
    view = make_view(application)
    index = view.model().index(0, 0)

    assert view.model().setData(index, '=1+1', Qt.EditRole)
    assert view.table[('A', datetime.datetime(2024, 1, 1))]['formula'].prop_value == '=1+1'
    assert view.need_update[-1].param_prop_name == 'formula'


def test_export_uses_model_values_without_eager_items(application, tmp_path):
    view = make_view(application)
    output = tmp_path / 'work-data.xlsx'

    view.export(str(output))
    worksheet = load_workbook(output).active

    assert worksheet.cell(1, 1).value == 'A'
    assert worksheet.cell(1, 2).value == 1.25
    assert view.model()._items == {}
