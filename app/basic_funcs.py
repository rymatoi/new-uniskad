import ast
import functools
import os
import time
from datetime import datetime

import xlrd
from PySide2.QtGui import QIcon
from PySide2.QtWidgets import QInputDialog, QLineEdit, QMessageBox
from PySide2.QtCore import QDir, Qt
from PySide2.QtWidgets import QFileDialog

from dateutil import parser
import re

from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.reader.excel import load_workbook
from openpyxl.workbook import Workbook

from app import app_logger

logger = app_logger.get_logger(__name__)


def parse_date(cell_value):
    # Define the desired date format
    date_format = '%Y-%m-%d %H:%M:%S'

    # Use regular expressions to clean up the string
    cleaned_value = re.sub(r'[^0-9\-:\s]', '', str(cell_value))

    # Insert a space between the date and time components if not present
    cleaned_value = re.sub(r'(\d{4}-\d{2}-\d{2})(\d{2}:\d{2}:\d{2})', r'\1 \2', cleaned_value)

    try:
        # Parse the cleaned string to a datetime object
        parsed_date = parser.parse(cleaned_value)

        # Format the datetime object to the desired format
        formatted_date = parsed_date.strftime(date_format)

        return formatted_date
    except ValueError:
        # Handle the case where parsing fails
        return None


def extract_time(cell_value):
    # Use regular expressions to extract time components
    time_match = re.search(r'(\d{2}:\d{2}:\d{2})', str(cell_value))

    if time_match:
        return time_match.group(1)
    else:
        return None


def extract_date(cell_value):
    # Define the desired date format
    date_format = '%d.%m.%Y'

    # Use regular expressions to clean up the string
    cleaned_value = re.sub(r'[^0-9\-:\s.]', '', str(cell_value))

    try:
        # Parse the cleaned string to a datetime object using datetime.strptime
        parsed_date = datetime.strptime(cleaned_value.strip(), '%d.%m.%Y')

        # Format the datetime object to extract the date
        formatted_date = parsed_date.strftime(date_format)

        return formatted_date
    except ValueError:
        # Handle the case where parsing fails
        error('Неправильный формат даты', 'Дата должна быть в формате "дд.мм.гггг"')
        return None


def get_text(title, label_text, default_text='', icon=":/uniskad.ico"):
    dialog = QInputDialog()
    dialog.setWindowIcon(QIcon(icon))
    dialog.setInputMode(QInputDialog.TextInput)
    dialog.setOkButtonText("Ок")
    dialog.setCancelButtonText("Отмена")
    dialog.setWindowTitle(title)
    dialog.setLabelText(label_text)
    dialog.setTextValue(default_text)
    dialog.setTextEchoMode(QLineEdit.Normal)
    dialog.setWindowFlags(dialog.windowFlags() & ~Qt.WindowContextHelpButtonHint)
    dialog.resize(300, 100)
    ok = dialog.exec_()
    name = dialog.textValue()
    if ok and name:
        return name


def get_files(caption, filter, single_selection=False):
    func = QFileDialog.getOpenFileName if single_selection else QFileDialog.getOpenFileNames
    selected_files = func(
        parent=None,
        caption=caption,
        dir=QDir.homePath(),
        filter=filter,
    )
    if single_selection:
        return selected_files[0]
    return [
        file for file in selected_files[0] if file
    ]


def get_answer(title, text):
    msg = QMessageBox(None)
    msg.setWindowFlags(msg.windowFlags() & ~Qt.WindowContextHelpButtonHint)
    msg.setWindowFlag(Qt.WindowStaysOnTopHint)
    msg.setWindowIcon(QIcon(":/uniskad.ico"))
    reply = msg.warning(None, title, text, QMessageBox.Yes | QMessageBox.No, )
    if reply == QMessageBox.Yes:
        return True
    return False


def error(title, text):
    msg = QMessageBox(None)
    msg.setWindowFlags(msg.windowFlags() & ~Qt.WindowContextHelpButtonHint)
    msg.setWindowFlag(Qt.WindowStaysOnTopHint)
    msg.setWindowIcon(QIcon(":/uniskad.ico"))
    reply = msg.critical(None, title, text)


def info(title, text):
    msg = QMessageBox(None)
    msg.setWindowFlags(msg.windowFlags() & ~Qt.WindowContextHelpButtonHint)
    msg.setWindowFlag(Qt.WindowStaysOnTopHint)
    msg.setWindowIcon(QIcon(":/uniskad.ico"))
    reply = msg.information(None, title, text)


def rgetattr(obj, attr, *args):
    def _getattr(obj, attr):
        return getattr(obj, attr, *args)

    return functools.reduce(_getattr, [obj] + attr.split('.'))


def export_file(filename="", caption="", filter="", parent=None) -> str:
    # Очищаем имя файла от недопустимых символов
    invalid_chars = '<>:"/\\|?*'
    clean_filename = ''.join(c for c in filename if c not in invalid_chars)

    path = os.path.join(QDir.homePath(), clean_filename)
    selected_file = QFileDialog.getSaveFileName(
        parent=parent, caption=caption, dir=path, filter=filter
    )
    return selected_file[0]


def check_recursion_limit(msg):
    def decorator(func):
        def wrapper(*args, **kwargs):
            try:
                result = func(*args, **kwargs)
            except RecursionError:
                result = msg
            return result

        return wrapper

    return decorator


def str_to_list(string):
    return ast.literal_eval(string)


def to_float(val):
    try:
        return float(str(val).replace(',', '.'))
    except Exception:
        return 0


def to_bool(val):
    bool_dict = {
        'True': True,
        'False': False,
        'true': True,
        'false': False
    }

    if isinstance(val, bool):
        return val

    if val in bool_dict:
        return bool_dict[val]
    else:
        return False


def xls2xlsx_(filename: str):
    if filename.endswith('.xls'):
        book = xlrd.open_workbook(filename)
        sh = book.sheet_by_index(0)
        wb = Workbook()
        ws = wb.active
        for row in range(sh.nrows):
            for column in range(sh.ncols):
                val = sh.cell_value(rowx=row, colx=column)
                if isinstance(val, str):
                    val = ILLEGAL_CHARACTERS_RE.sub(r'', val)
                ws.cell(row + 1, column + 1).value = val
    else:
        wb = load_workbook(filename)
    return wb


def float_to_excel(val):
    return str(val).replace('.', ',')


def excel_to_float(val):
    return str(val).replace(',', '.')


def timing_decorator(func):
    @functools.wraps(func)
    def wrapper(*args, **kwargs):
        start_time = time.perf_counter()
        result = func(*args, **kwargs)
        logger.debug(
            "Function %s.%s completed in %.4f seconds",
            func.__module__,
            func.__qualname__,
            time.perf_counter() - start_time,
        )
        return result

    return wrapper
