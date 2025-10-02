from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.drawing.colors import ColorChoice
from openpyxl.chart.marker import Marker
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from PySide2.QtGui import QColor
from PySide2.QtCore import Qt

from openpyxl.chart.axis import ChartLines
from app.plugins.project.core.constants import GraphConstants


class PlotExcelExporter:
    """Класс для экспорта данных графика в Excel с сохранением стилей."""

    # Маппинг типов линий Qt в типы линий Excel
    LINE_STYLE_MAP = {
        Qt.NoPen: 'none',
        Qt.SolidLine: 'solid',
        Qt.DashLine: 'dash',
        Qt.DotLine: 'dot',
        Qt.DashDotLine: 'dashDot',
        Qt.DashDotDotLine: 'sysDashDotDot',
    }

    def __init__(self):
        self.wb = Workbook()
        # Создаем два листа: для данных и для графика
        self.ws_data = self.wb.active
        self.ws_data.title = "Данные"
        self.ws_plot = self.wb.create_sheet(title="График")
        self.current_row = 1
        self.data_ranges = []  # Для хранения диапазонов данных каждой кривой

    def export(self, plot_view, filename):
        """Экспортирует данные и график в Excel файл."""
        self._export_curves(plot_view)
        self._create_plot(plot_view)
        self._auto_adjust_columns()
        self.wb.save(filename)

    @staticmethod
    def _normalize_value(value):
        """Приводит значения numpy.* к python типам для корректной записи в Excel."""
        if hasattr(value, 'item'):
            try:
                return value.item()
            except Exception:
                pass
        return value

    def _ensure_qcolor(self, value):
        """Возвращает валидный QColor для произвольного входного значения."""
        if isinstance(value, QColor):
            return value
        if value is None:
            return QColor('#000000')
        if isinstance(value, (tuple, list)) and len(value) >= 3:
            return QColor(*value[:3])
        color = QColor(value)
        if color.isValid():
            return color
        return QColor('#000000')

    def _get_excel_color(self, value):
        """Преобразует цвет в формат Excel."""
        qt_color = self._ensure_qcolor(value)
        return f"{qt_color.red():02x}{qt_color.green():02x}{qt_color.blue():02x}"

    @staticmethod
    def _extract_curve_points(curve):
        """Возвращает списки X и Y для кривой или эпюры."""
        if hasattr(curve, 'xData') and curve.xData is not None:
            return list(curve.xData), list(curve.yData)

        if hasattr(curve, 'scatter') and curve.scatter is not None:
            points = curve.scatter.points()
            return [p.pos().x() for p in points], [p.pos().y() for p in points]

        return [], []

    def _qt_pen_to_excel_style(self, line_style):
        """Преобразует Qt стиль линии в стиль Excel."""
        pen_style = GraphConstants.resolve_pen_style(line_style)
        return self.LINE_STYLE_MAP.get(pen_style, 'solid')

    def _export_curves(self, plot_view):
        """Экспортирует данные кривых и сохраняет их диапазоны."""
        for curve in plot_view.curve_items:
            if not curve.isVisible():
                continue
            x_values, y_values = self._extract_curve_points(curve)
            if not x_values or not y_values:
                continue

            style_config = getattr(curve, 'style_config', getattr(curve, '_style_config', GraphConstants.DEFAULT_STYLE))
            style = dict(style_config)

            curve_name = curve.name() if hasattr(curve, 'name') else style.get('name', 'Curve')
            rgb_color = self._get_excel_color(style.get('color', '#000000'))

            title_fill = PatternFill(start_color=rgb_color, end_color=rgb_color, fill_type="solid")
            border = Border(
                left=Side(style='thin', color=rgb_color),
                right=Side(style='thin', color=rgb_color),
                top=Side(style='thin', color=rgb_color),
                bottom=Side(style='thin', color=rgb_color)
            )

            title_row = self.current_row
            self.ws_data.merge_cells(f'A{title_row}:D{title_row}')
            title_cell = self.ws_data.cell(row=title_row, column=1, value=curve_name)
            title_cell.fill = title_fill
            title_cell.font = Font(color="FFFFFF", bold=True)
            title_cell.alignment = Alignment(horizontal='left')

            header_row = title_row + 1
            headers = ['№ точки', 'X', 'Y', 'Статус']
            for col, header in enumerate(headers, 1):
                header_cell = self.ws_data.cell(row=header_row, column=col, value=header)
                header_cell.font = Font(bold=True)
                header_cell.border = border

            row_pointer = header_row
            selected_indices = sorted(plot_view.selected_points.get(curve, set()))
            selected_set = {idx for idx in selected_indices if 0 <= idx < len(x_values)}
            selected_points = []

            for idx, (x_val, y_val) in enumerate(zip(x_values, y_values), start=1):
                row_pointer += 1
                status = "Выделена" if (idx - 1) in selected_set else ""
                row_data = [
                    idx,
                    self._normalize_value(x_val),
                    self._normalize_value(y_val),
                    status
                ]
                for col, value in enumerate(row_data, 1):
                    data_cell = self.ws_data.cell(row=row_pointer, column=col, value=value)
                    data_cell.border = border
                    if status and col == 4:
                        data_cell.font = Font(color="FF0000", bold=True)

                if status:
                    selected_points.append((
                        self._normalize_value(x_values[idx - 1]),
                        self._normalize_value(y_values[idx - 1])
                    ))

            data_start_row = header_row + 1
            data_end_row = row_pointer if x_values else header_row

            if hasattr(curve, 'get_point_info'):
                for point_index in range(len(x_values)):
                    additional_info = curve.get_point_info(point_index)
                    if additional_info:
                        for key, value in additional_info.items():
                            row_pointer += 1
                            key_cell = self.ws_data.cell(row=row_pointer, column=1, value=key)
                            key_cell.font = Font(bold=True)
                            self.ws_data.cell(row=row_pointer, column=2, value=str(value))

            self.data_ranges.append({
                'name': curve_name,
                'data_start_row': data_start_row,
                'data_end_row': data_end_row,
                'style': style,
                'symbol': style.get('symbol'),
                'selected_points': selected_points,
                'point_size': getattr(curve, 'point_size', style.get('symbol_size', GraphConstants.DEFAULT_STYLE['symbol_size'])),
                'selected_factor': getattr(curve, 'selected_curve_point_size', 0.5),
                'visible': curve.isVisible()
            })

            self.current_row = row_pointer + 2

    def _get_marker_style(self, symbol):
        """Определяет стиль маркера точки для Excel."""
        symbol_map = {
            'o': 'circle',
            's': 'square',
            't': 'triangle',
            't1': 'triangle',
            't2': 'triangle',
            't3': 'triangle',
            'd': 'diamond',
            'p': 'diamond',
            'h': 'diamond',
            '+': 'plus',
            'x': 'x',
            'star': 'star',
            None: 'none'
        }
        return symbol_map.get(symbol, 'none')

    def _create_plot(self, plot_view):
        """Создает график в Excel."""
        chart = ScatterChart()
        chart.title = None
        chart.x_axis.title = plot_view.plotItem.axes['bottom']['item'].label.toPlainText()
        chart.y_axis.title = plot_view.plotItem.axes['left']['item'].label.toPlainText()

        # Настраиваем сетку
        chart.x_axis.majorGridlines = ChartLines()
        chart.y_axis.majorGridlines = ChartLines()

        # Делаем сетку серой и тонкой
        chart.x_axis.majorGridlines.spPr = GraphicalProperties(
            ln=LineProperties(
                solidFill=ColorChoice(srgbClr="CCCCCC"),  # Светло-серый цвет
            )
        )
        chart.x_axis.majorGridlines.spPr.ln.w = int(0.5 * 9525)  # Тонкая линия

        chart.y_axis.majorGridlines.spPr = GraphicalProperties(
            ln=LineProperties(
                solidFill=ColorChoice(srgbClr="CCCCCC"),
            )
        )
        chart.y_axis.majorGridlines.spPr.ln.w = int(0.5 * 9525)

        extra_row_pointer = max(self.current_row, 1) + 1

        for curve_data in self.data_ranges:
            if not curve_data['visible']:
                continue

            data_start = curve_data['data_start_row']
            data_end = curve_data['data_end_row']
            if data_end < data_start:
                continue

            x_ref = Reference(self.ws_data, min_col=2, min_row=data_start, max_row=data_end)
            y_ref = Reference(self.ws_data, min_col=3, min_row=data_start, max_row=data_end)
            series = Series(y_ref, x_ref, title=curve_data['name'])

            style = curve_data['style']
            excel_line_style = self._qt_pen_to_excel_style(style.get('line_style'))
            try:
                line_width = float(style.get('width', GraphConstants.DEFAULT_STYLE['width']))
            except (TypeError, ValueError):
                line_width = float(GraphConstants.DEFAULT_STYLE['width'])
            show_line = excel_line_style != 'none' and line_width > 0

            marker_symbol = self._get_marker_style(curve_data['symbol'])
            if marker_symbol != 'none':
                try:
                    marker_size = int(curve_data.get('point_size', style.get('symbol_size', GraphConstants.DEFAULT_STYLE['symbol_size'])))
                except (TypeError, ValueError):
                    marker_size = int(GraphConstants.DEFAULT_STYLE['symbol_size'])
                marker_size = max(marker_size, 1)
                fill_value = style.get('fill_color')
                if fill_value is None:
                    fill_value = style.get('color')
                border_value = style.get('symbol_color')
                if border_value is None:
                    border_value = style.get('color')
                marker_fill_color = self._get_excel_color(fill_value)
                marker_border_color = self._get_excel_color(border_value)
                series.marker = Marker(
                    symbol=marker_symbol,
                    size=marker_size,
                    spPr=GraphicalProperties(
                        solidFill=ColorChoice(srgbClr=marker_fill_color),
                        ln=LineProperties(solidFill=ColorChoice(srgbClr=marker_border_color))
                    )
                )
            else:
                series.marker = Marker(symbol='none')

            if not show_line:
                series.graphicalProperties.line = LineProperties()
                series.graphicalProperties.line.noFill = True
            else:
                line_color = self._get_excel_color(style.get('color', '#000000'))
                line_props = LineProperties(solidFill=ColorChoice(srgbClr=line_color))
                if excel_line_style != 'solid':
                    line_props.prstDash = excel_line_style
                line_props.w = int(line_width * 9525)
                series.graphicalProperties.line = line_props

            if curve_data['selected_points']:
                sel_start, sel_end, extra_row_pointer = self._write_selected_points(curve_data['selected_points'], extra_row_pointer)
                sel_x_ref = Reference(self.ws_data, min_col=6, min_row=sel_start, max_row=sel_end)
                sel_y_ref = Reference(self.ws_data, min_col=7, min_row=sel_start, max_row=sel_end)
                selected_series = Series(sel_y_ref, sel_x_ref, title=f"Выделенные точки ({curve_data['name']})")

                try:
                    base_marker = int(curve_data.get('point_size', style.get('symbol_size', GraphConstants.DEFAULT_STYLE['symbol_size'])))
                except (TypeError, ValueError):
                    base_marker = int(GraphConstants.DEFAULT_STYLE['symbol_size'])
                base_marker = max(base_marker, 1)
                factor = curve_data.get('selected_factor', 0.5)
                try:
                    factor = float(factor)
                except (TypeError, ValueError):
                    factor = 0.5
                highlight_size = max(int(base_marker * (1 + factor)), base_marker + 1)
                highlight_color = ColorChoice(srgbClr="FF0000")
                highlight_symbol = marker_symbol if marker_symbol != 'none' else 'circle'

                selected_series.marker = Marker(
                    symbol=highlight_symbol,
                    size=highlight_size,
                    spPr=GraphicalProperties(
                        solidFill=highlight_color,
                        ln=LineProperties(solidFill=highlight_color)
                    )
                )
                selected_series.graphicalProperties.line = LineProperties()
                selected_series.graphicalProperties.line.noFill = True

                chart.series.append(selected_series)

            chart.series.append(series)

        self.ws_plot.add_chart(chart, "B2")
        chart.height = 20
        chart.width = 30

        self.current_row = extra_row_pointer

    def _write_selected_points(self, points, start_row):
        """Записывает координаты выделенных точек в отдельные столбцы."""
        if not points:
            return None, None, start_row

        start_row = max(int(start_row), 1)
        for offset, (x_val, y_val) in enumerate(points):
            row_idx = start_row + offset
            self.ws_data.cell(row=row_idx, column=6, value=self._normalize_value(x_val))
            self.ws_data.cell(row=row_idx, column=7, value=self._normalize_value(y_val))

        end_row = start_row + len(points) - 1
        return start_row, end_row, end_row + 2

    def _auto_adjust_columns(self):
        """Автоматически подбирает ширину столбцов в листе данных."""
        for column in self.ws_data.columns:
            max_length = 0
            column_cells = list(column)
            for cell in column_cells:
                value = cell.value
                if value is None:
                    continue
                length = len(str(value))
                if length > max_length:
                    max_length = length
            adjusted_width = max_length + 2 if max_length > 0 else 8
            self.ws_data.column_dimensions[get_column_letter(column_cells[0].column)].width = adjusted_width
