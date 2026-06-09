import time

from openpyxl.workbook import Workbook
from PySide2.QtCore import QAbstractTableModel, QLocale, Qt
from PySide2.QtGui import QBrush, QColor, QFont
from PySide2.QtWidgets import QAbstractItemView, QTableView

from app import app_logger
from app.plugins.base_state.widgets import TableWidget, TableItem
from db import sp
from db.tables import IMPORT_FILE_DATA

logger = app_logger.get_logger(__name__)


class WorkDataTableItem(TableItem):
    def update_cell(self, prop, value):
        tmp = self.get_prop_template(prop, value)
        self.tableWidget().need_update.append(tmp)
        self.add_prop(tmp)


class WorkDataTableWidget(TableWidget):
    TABLE_ITEM = WorkDataTableItem

    def get_update_cells(self):
        return [cell.table_fit(IMPORT_FILE_DATA) for cell in self.need_update]

    def load_table(self, db_objects):
        started = time.perf_counter()
        super().load_table(db_objects)
        logger.debug("WorkDataTableWidget: loaded QTableWidget fallback in %.4f seconds",
                     time.perf_counter() - started)

    def update_table(self):
        success = sp.new_upd_excel_data_array(self.get_update_cells())
        if success:
            self.need_update = []


# The model/view path intentionally lives beside the old widget implementation.
# It can therefore be enabled for large, read-only Work Data views without
# changing the behavior of other table users.


class WorkDataTableModel(QAbstractTableModel):
    """Lazy Work Data model that creates no visual object per table cell."""

    def __init__(self, parent=None):
        super().__init__(parent)
        self.rows = {}
        self.columns = {}
        self.table = {}
        self.ord_rows = []
        self.ord_columns = []
        self._row_indexes = {}
        self._column_indexes = {}

    def load_data(self, db_objects):
        started = time.perf_counter()
        self.beginResetModel()
        self.rows = {}
        self.columns = {}
        self.table = {}

        for obj in db_objects:
            param_name = obj.excel_param_name if obj.sprav_name is None else obj.sprav_name
            if obj.date_time_izm is None:
                target = self.rows.setdefault((param_name, None), {})
            elif param_name is None:
                target = self.columns.setdefault((None, obj.date_time_izm), {})
            else:
                target = self.table.setdefault((param_name, obj.date_time_izm), {})
            target[obj.prop_name] = obj

        self.ord_columns = [
            column[1] for column in sorted(
                self.columns, key=lambda column: int(self.columns[column]['column_npp'].prop_value))
        ]
        self.ord_rows = [
            row[0] for row in sorted(
                self.rows, key=lambda row: self._row_prop(row, 'row_npp', int, None))
        ]
        self._row_indexes = {key: index for index, key in enumerate(self.ord_rows)}
        self._column_indexes = {key: index for index, key in enumerate(self.ord_columns)}
        self.endResetModel()
        logger.debug(
            "WorkDataTableModel: prepared model data in %.4f seconds "
            "(%d records, %d rows, %d columns, %d populated cells)",
            time.perf_counter() - started,
            len(db_objects),
            len(self.ord_rows),
            len(self.ord_columns),
            len(self.table),
        )

    def rowCount(self, parent=None):
        return 0 if parent is not None and parent.isValid() else len(self.ord_rows)

    def columnCount(self, parent=None):
        return 0 if parent is not None and parent.isValid() else len(self.ord_columns)

    def data(self, index, role=Qt.DisplayRole):
        if not index.isValid():
            return None
        cell = self._cell(index.row(), index.column())
        if role == Qt.DisplayRole:
            return self._display_value(cell, self.ord_rows[index.row()])
        if role == Qt.EditRole:
            return self._get(cell, 'formula', str, '')
        if role == Qt.BackgroundRole:
            if self._get(cell, 'broken', bool, False):
                return QBrush(Qt.lightGray)
            color = self._get(cell, 'font_bgcolor', str, None)
            return QColor(color) if color else None
        if role == Qt.ForegroundRole:
            color = self._get(cell, 'font_text_color', str, None)
            return QColor(color) if color else None
        if role == Qt.FontRole:
            font = QFont()
            font.setBold(self._get(cell, 'font_bold', bool, False))
            font.setItalic(self._get(cell, 'font_italic', bool, False))
            size = self._get(cell, 'font_size', int, None)
            if size:
                font.setPixelSize(size)
            return font
        return None

    def setData(self, index, value, role=Qt.EditRole):
        # Formula editing remains on the old QTableWidget fallback until its
        # dependency graph can be moved out of TableItem safely.
        return False

    def flags(self, index):
        if not index.isValid():
            return Qt.NoItemFlags
        return Qt.ItemIsEnabled | Qt.ItemIsSelectable

    def headerData(self, section, orientation, role=Qt.DisplayRole):
        if role != Qt.DisplayRole:
            return None
        values = self.ord_rows if orientation == Qt.Vertical else self.ord_columns
        return values[section] if 0 <= section < len(values) else None

    def _cell(self, row, column):
        return self.table.get((self.ord_rows[row], self.ord_columns[column]), {})

    @staticmethod
    def _get(cell, prop, cast_type=None, default=None):
        obj = cell.get(prop)
        if obj is None or not obj.prop_value:
            return default
        value = obj.prop_value
        if cast_type is bool:
            return value if isinstance(value, bool) else str(value).lower() == 'true'
        if cast_type is not None:
            try:
                return cast_type(value)
            except (TypeError, ValueError):
                return default
        return value

    def _row_prop(self, row, prop, cast_type, default=None):
        return self._get(self.rows.get(row, {}), prop, cast_type, default)

    def _display_value(self, cell, row):
        accuracy = self._row_prop((row, None), 'accuracy', int, 2)
        plus_value = self._row_prop((row, None), 'plus_value', float, 0)
        mul_value = self._row_prop((row, None), 'mul_value', float, 1)
        value = self._get(cell, 'cformula', str, None)
        if value is None:
            value = self._get(cell, 'value', str, None)
        try:
            numeric_value = float(value)
        except (TypeError, ValueError):
            return value
        numeric_value = numeric_value + plus_value if plus_value else numeric_value * mul_value
        return QLocale().toString(numeric_value, 'f', accuracy)

    def is_secret_row(self, row):
        return self._row_prop((self.ord_rows[row], None), 'is_secret', bool, False)

    def export(self, filename):
        workbook = Workbook()
        worksheet = workbook.active
        for row, row_name in enumerate(self.ord_rows):
            worksheet.cell(row + 1, 1).value = row_name
            for column in range(len(self.ord_columns)):
                cell = self._cell(row, column)
                value = self._get(cell, 'cformula', str, self._get(cell, 'value', str, ''))
                try:
                    value = float(value)
                except (TypeError, ValueError):
                    pass
                worksheet.cell(row + 1, column + 2).value = value
        workbook.save(filename)


class WorkDataTableView(QTableView):
    """Opt-in, read-only model/view replacement for WorkDataTableWidget."""

    def __init__(self, parent=None, main_window=None):
        super().__init__(parent)
        self._parent = parent
        self.mw = main_window
        self.setLocale(QLocale(QLocale.English, QLocale.UnitedKingdom))
        self.setSelectionMode(QAbstractItemView.ExtendedSelection)
        self.setSelectionBehavior(QAbstractItemView.SelectItems)
        self.setModel(WorkDataTableModel(self))

    @property
    def rows(self):
        return self.model().rows

    @property
    def columns(self):
        return self.model().columns

    @property
    def table(self):
        return self.model().table

    @property
    def ord_rows(self):
        return self.model().ord_rows

    @property
    def ord_columns(self):
        return self.model().ord_columns

    def load_table(self, db_objects):
        started = time.perf_counter()
        self.model().load_data(db_objects)
        if sp.get_session_role_secret_grantness() is False:
            for row in range(self.model().rowCount()):
                self.setRowHidden(row, self.model().is_secret_row(row))
        logger.debug("WorkDataTableView: set up QTableView in %.4f seconds", time.perf_counter() - started)

    def export(self, filename):
        self.model().export(filename)

    def update_table(self):
        logger.warning("Work Data model/view mode is read-only; no changes were saved")

    def apply_filters(self, filters):
        logger.warning("Work Data model/view mode does not yet support saved broken-value filters")

    def update_row_obj(self, row, prop_name, obj):
        key = row if isinstance(row, tuple) else (row, None)
        self.model().rows.setdefault(key, {})[prop_name] = obj

    def update_column_obj(self, column, prop_name, obj):
        key = column if isinstance(column, tuple) else (None, column)
        self.model().columns.setdefault(key, {})[prop_name] = obj
