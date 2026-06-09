import ast
import time
from copy import copy
from collections.abc import Iterable

from openpyxl.workbook import Workbook
from PySide2.QtCore import QAbstractTableModel, QLocale, QTimer, Qt, Signal
from PySide2.QtGui import QKeySequence
from PySide2.QtWidgets import (QAbstractItemView, QApplication, QInputDialog,
                               QTableView, QTableWidgetSelectionRange)

from app import app_logger
from app.plugins.base_state.widgets import TableItem
from db import sp

logger = app_logger.get_logger(__name__)

def clone_property_record(template, prop_name, prop_value, key=None):
    """Create an independent property record from a compatible record template."""
    record = copy(template)
    record.id_record = None
    record.param_prop_name = str(prop_name)
    record.prop_name = str(prop_name)
    record.prop_value = str(prop_value)
    record.deleted = False
    if key is not None:
        record.excel_param_name, record.date_time_izm = key
        record.sprav_name = None
    return record


class LazyTableModelItem(TableItem):
    """Lazily-created compatibility item for business rules still owned by TableItem.

    These objects are only made for cells that code or the viewport touches; the
    model never eagerly creates one object per table cell.
    """

    def __init__(self, table, cell, key):
        self._table = table
        super().__init__(cell, key)

    def tableWidget(self):
        return self._table

    def row(self):
        return self._table.model()._row_indexes.get(self.key[0], -1)

    def column(self):
        return self._table.model()._column_indexes.get(self.key[1], -1)

    def text(self):
        value = self.value()
        return '' if value is None else str(value)

    def update_cell(self, prop, value):
        # A mutation of a malformed legacy/sparse cell must first restore the
        # base records expected by TableItem business rules.
        model = self.tableWidget().model()
        model.item(model.index(self.row(), self.column()), create=True)
        tmp = self.get_prop_template(prop, value)
        self.tableWidget().need_update.append(tmp)
        self.add_prop(tmp)
        index = self._table.indexFromItem(self)
        if index.isValid():
            self._table.model().dataChanged.emit(
                index, index,
                [Qt.DisplayRole, Qt.EditRole, Qt.BackgroundRole, Qt.ForegroundRole,
                 Qt.FontRole, Qt.TextAlignmentRole])
            self._table.itemChanged.emit(self)


class LazyTableModel(QAbstractTableModel):
    """Reusable lazy record-backed model preserving legacy TableItem rules."""

    ITEM_CLASS = LazyTableModelItem

    def __init__(self, parent=None):
        started = time.perf_counter()
        super().__init__(parent)
        self.rows = {}
        self.columns = {}
        self.table = {}
        self.ord_rows = []
        self.ord_columns = []
        self._row_indexes = {}
        self._column_indexes = {}
        self._items = {}
        logger.info("LazyTableModel: initialized in %.4f seconds", time.perf_counter() - started)

    def load_data(self, db_objects):
        started = time.perf_counter()
        self.beginResetModel()
        self.rows = {}
        self.columns = {}
        self.table = {}
        self._items = {}

        for obj in db_objects:
            self._object_for(obj).update({obj.prop_name: obj})

        self.ord_columns = [
            column[1] for column in sorted(
                self.columns, key=lambda column: int(self.columns[column]['column_npp'].prop_value))
        ]
        self.ord_rows = [
            row[0] for row in sorted(
                self.rows, key=lambda row: self._row_prop(row, 'row_npp', int, None))
        ]
        self._rebuild_indexes()
        self.endResetModel()
        logger.info(
            "LazyTableModel: prepared model data in %.4f seconds "
            "(%d records, %d rows, %d columns, %d populated cells)",
            time.perf_counter() - started, len(db_objects), len(self.ord_rows),
            len(self.ord_columns), len(self.table))

    def _object_for(self, obj):
        param_name = obj.excel_param_name if obj.sprav_name is None else obj.sprav_name
        if obj.date_time_izm is None:
            return self.rows.setdefault((param_name, None), {})
        if param_name is None:
            return self.columns.setdefault((None, obj.date_time_izm), {})
        return self.table.setdefault((param_name, obj.date_time_izm), {})

    def _rebuild_indexes(self):
        self._row_indexes = {key: index for index, key in enumerate(self.ord_rows)}
        self._column_indexes = {key: index for index, key in enumerate(self.ord_columns)}

    def rowCount(self, parent=None):
        return 0 if parent is not None and parent.isValid() else len(self.ord_rows)

    def columnCount(self, parent=None):
        return 0 if parent is not None and parent.isValid() else len(self.ord_columns)

    def data(self, index, role=Qt.DisplayRole):
        if not index.isValid():
            return None
        item = self.item(index)
        if item is None:
            return '' if role in (Qt.DisplayRole, Qt.EditRole) else None
        return item.data(role)

    def setData(self, index, value, role=Qt.EditRole):
        if role != Qt.EditRole or not index.isValid():
            return False
        value = '' if value is None else str(value)
        item = self.item(index, create=True)
        if item is None or value == item.get('formula', str, ''):
            return False
        item.setData(role, value)
        self.dataChanged.emit(index, index, [Qt.DisplayRole, Qt.EditRole])
        table = self.parent()
        if table is not None:
            table.itemChanged.emit(item)
        return True

    def flags(self, index):
        if not index.isValid():
            return Qt.NoItemFlags
        return Qt.ItemIsEnabled | Qt.ItemIsSelectable | Qt.ItemIsEditable

    def headerData(self, section, orientation, role=Qt.DisplayRole):
        if role != Qt.DisplayRole:
            return None
        if orientation == Qt.Vertical:
            return self.ord_rows[section] if 0 <= section < len(self.ord_rows) else None
        # QTableWidget's default horizontal headers are one-based section
        # numbers; retain that legacy display instead of exposing timestamps.
        return section + 1 if 0 <= section < len(self.ord_columns) else None

    def item(self, index, create=False):
        if not index.isValid() or not (0 <= index.row() < len(self.ord_rows)) or not (
                0 <= index.column() < len(self.ord_columns)):
            return None
        key = (self.ord_rows[index.row()], self.ord_columns[index.column()])
        if create:
            self.ensure_cell(key)
        cell = self.table.get(key)
        if cell is None:
            return None
        item = self._items.get(key)
        if item is None:
            item = self.ITEM_CLASS(self.parent(), cell, key)
            self._items[key] = item
        return item

    def ensure_cell(self, key):
        """Materialize a complete independent cell and queue newly-created records."""
        cell = self.table.get(key)
        template = next(iter(cell.values()), None) if cell is not None else None
        for records_by_key in (self.table, self.rows, self.columns):
            if template is not None:
                break
            template = next((record for properties in records_by_key.values()
                             for record in properties.values()), None)
        if template is None:
            return None

        cell = self.table.setdefault(key, {})
        for prop_name, prop_value in (('value', '0'), ('type', 'cell')):
            if prop_name in cell:
                continue
            record = clone_property_record(template, prop_name, prop_value, key)
            cell[prop_name] = record
            self.parent().need_update.append(record)
            template = record
        return cell

    @staticmethod
    def _get(cell, prop, cast_type=None, default=None):
        obj = cell.get(prop)
        if obj is None or not obj.prop_value:
            return default
        value = obj.prop_value
        if cast_type is bool:
            return value if isinstance(value, bool) else str(value).lower() == 'true'
        if cast_type is not None:
            try:
                return cast_type(value)
            except (TypeError, ValueError):
                return default
        return value

    def _row_prop(self, row, prop, cast_type, default=None):
        return self._get(self.rows.get(row, {}), prop, cast_type, default)

    def is_secret_row(self, row):
        return self._row_prop((self.ord_rows[row], None), 'is_secret', bool, False)

    def export(self, filename):
        workbook = Workbook()
        worksheet = workbook.active
        for row, row_name in enumerate(self.ord_rows):
            worksheet.cell(row + 1, 1).value = row_name
            for column in range(len(self.ord_columns)):
                cell = self.table.get((row_name, self.ord_columns[column]), {})
                value = self._get(cell, 'cformula', str, self._get(cell, 'value', str, ''))
                try:
                    value = float(value)
                except (TypeError, ValueError):
                    pass
                worksheet.cell(row + 1, column + 2).value = value
        workbook.save(filename)


class ModelViewTable(QTableView):
    """Reusable QTableView exposing the legacy QTableWidget-facing API."""

    MODEL_CLASS = LazyTableModel
    TABLE_FIT = None
    SAVE_FUNCTION = None

    itemChanged = Signal(object)

    def __init__(self, parent=None, main_window=None):
        started = time.perf_counter()
        super().__init__(parent)
        self._parent = parent
        self.mw = main_window
        self.search_string = None
        self.need_update = []
        self._highlighted_items = []
        self.setLocale(QLocale(QLocale.English, QLocale.UnitedKingdom))
        self.setSelectionMode(QAbstractItemView.ExtendedSelection)
        self.setSelectionBehavior(QAbstractItemView.SelectItems)
        self.setEditTriggers(QAbstractItemView.DoubleClicked | QAbstractItemView.EditKeyPressed |
                             QAbstractItemView.AnyKeyPressed)
        self.setModel(self.MODEL_CLASS(self))
        logger.info("ModelViewTable: base setup completed in %.4f seconds",
                    time.perf_counter() - started)

    @property
    def rows(self):
        return self.model().rows

    @property
    def columns(self):
        return self.model().columns

    @property
    def table(self):
        return self.model().table

    @property
    def ord_rows(self):
        return self.model().ord_rows

    @property
    def ord_columns(self):
        return self.model().ord_columns

    def load_table(self, db_objects):
        started = time.perf_counter()
        self.model().load_data(db_objects)
        if sp.get_session_role_secret_grantness() is False:
            self.hide_secret_rows()
        self.clear_highlights()
        logger.info("ModelViewTable: loaded and set up QTableView in %.4f seconds",
                    time.perf_counter() - started)

    def item(self, row, column, create=False):
        return self.model().item(self.model().index(row, column), create=create)

    def itemFromIndex(self, index, create=False):
        return self.model().item(index, create=create)

    def ensureItem(self, index):
        return self.model().item(index, create=True)

    def indexFromItem(self, item):
        if item is None:
            return self.model().index(-1, -1)
        return self.model().index(self.model()._row_indexes.get(item.key[0], -1),
                                  self.model()._column_indexes.get(item.key[1], -1))

    def currentItem(self):
        return self.itemFromIndex(self.currentIndex())

    def currentRow(self):
        return self.currentIndex().row()

    def currentColumn(self):
        return self.currentIndex().column()

    def selectedItems(self):
        return [self.itemFromIndex(index) for index in self.selectedIndexes()]

    def selectedRanges(self):
        return [QTableWidgetSelectionRange(selection.top(), selection.left(),
                                           selection.bottom(), selection.right())
                for selection in self.selectionModel().selection()]

    def rowCount(self):
        return self.model().rowCount()

    def columnCount(self):
        return self.model().columnCount()

    def setCurrentCell(self, row, column):
        self.setCurrentIndex(self.model().index(row, column))

    def get_update_cells(self):
        return [cell.table_fit(self.TABLE_FIT) for cell in self.need_update]

    def update_table(self):
        started = time.perf_counter()
        success = self.SAVE_FUNCTION(self.get_update_cells())
        logger.info("%s: save/update path completed in %.4f seconds (%d queued records)",
                    type(self).__name__, time.perf_counter() - started, len(self.need_update))
        if success:
            self.need_update = []
            self.post_save()
        return success

    def post_save(self):
        """Plugin-specific hook invoked after a successful save."""

    def export(self, filename):
        self.model().export(filename)

    def get_row_prop(self, row, prop_name, type_, default=None):
        key = row if isinstance(row, tuple) else (row, None)
        return self.model()._get(self.rows.get(key, {}), prop_name, type_, default)

    def get_column_prop(self, column, prop_name, type_, default=None):
        key = column if isinstance(column, tuple) else (None, column)
        return self.model()._get(self.columns.get(key, {}), prop_name, type_, default)

    def get_row_npp(self, row):
        return self.get_row_prop(row, 'row_npp', int, None)

    def get_column_npp(self, column):
        return self.get_column_prop(column, 'column_npp', int, None)

    def update_row_obj(self, row, prop_name, obj):
        key = row if isinstance(row, tuple) else (row, None)
        self.rows.setdefault(key, {})[prop_name] = obj
        row_index = self.model()._row_indexes.get(key[0])
        if row_index is not None and self.model().columnCount():
            self.model().dataChanged.emit(self.model().index(row_index, 0),
                                          self.model().index(row_index, self.model().columnCount() - 1))

    def update_column_obj(self, column, prop_name, obj):
        key = column if isinstance(column, tuple) else (None, column)
        self.columns.setdefault(key, {})[prop_name] = obj
        column_index = self.model()._column_indexes.get(key[1])
        if column_index is not None and self.model().rowCount():
            self.model().dataChanged.emit(self.model().index(0, column_index),
                                          self.model().index(self.model().rowCount() - 1, column_index))

    def update_ord_row(self, new_name, old_name):
        row_index = self.model()._row_indexes.get(old_name)
        if row_index is None or new_name == old_name:
            return
        self.model().beginResetModel()
        self.ord_rows[row_index] = new_name
        row_props = self.rows.pop((old_name, None), None)
        if row_props is not None:
            self.rows[(new_name, None)] = row_props
            for prop in row_props.values():
                prop.excel_param_name = new_name
        for column in self.ord_columns:
            cell = self.table.pop((old_name, column), None)
            if cell is None:
                continue
            self.table[(new_name, column)] = cell
            for prop in cell.values():
                prop.excel_param_name = new_name
        for cell in self.need_update:
            if cell.excel_param_name == old_name:
                cell.excel_param_name = new_name
        self.model()._items = {}
        self.model()._rebuild_indexes()
        self.model().endResetModel()

    def add_row(self, db_objects, position=None):
        if not db_objects:
            return
        new_row = db_objects[0].excel_param_name
        self.model().beginResetModel()
        for obj in db_objects:
            self.model()._object_for(obj).update({obj.prop_name: obj})
        if position is None or position < 0 or position > len(self.ord_rows):
            position = len(self.ord_rows)
        if new_row not in self.ord_rows:
            self.ord_rows.insert(position, new_row)
        self.model()._items = {}
        self.model()._rebuild_indexes()
        self.model().endResetModel()
        row_index = self.model()._row_indexes.get(new_row)
        if row_index is not None:
            for column in range(len(self.ord_columns)):
                item = self.item(row_index, column)
                if item is not None and item.get('formula', str, None):
                    item.calculate_formula()

    def add_column(self, db_objects):
        if not db_objects:
            return
        new_column = db_objects[0].date_time_izm
        self.model().beginResetModel()
        for obj in db_objects:
            self.model()._object_for(obj).update({obj.prop_name: obj})
        if new_column not in self.ord_columns:
            self.ord_columns.append(new_column)
        self.model()._items = {}
        self.model()._rebuild_indexes()
        self.model().endResetModel()

    def _queue_deleted_records(self, records, seen):
        """Mark existing records deleted and retain them for the next save."""
        queued_ids = {id(record) for record in self.need_update}
        marked = 0
        for record in records:
            record_id = id(record)
            if record_id in seen:
                continue
            seen.add(record_id)
            record.deleted = True
            if record_id not in queued_ids:
                self.need_update.append(record)
                queued_ids.add(record_id)
            marked += 1
        return marked

    def queue_row_deletion(self, row_key):
        """Queue all existing sparse cell and metadata records for a row."""
        before = len(self.need_update)
        seen = set()
        cell_records = 0
        property_records = 0

        for (record_row, record_column), properties in list(self.table.items()):
            matching = (properties.values() if record_row == row_key else
                        (record for record in properties.values()
                         if getattr(record, 'excel_param_name', None) == row_key))
            marked = self._queue_deleted_records(matching, seen)
            if record_column is None or record_column not in self.ord_columns:
                property_records += marked
            else:
                cell_records += marked

        for (record_row, _), properties in list(self.rows.items()):
            matching = (properties.values() if record_row == row_key else
                        (record for record in properties.values()
                         if getattr(record, 'excel_param_name', None) == row_key))
            property_records += self._queue_deleted_records(matching, seen)

        # Include unusual metadata records without assuming where a plugin stores them.
        for properties in list(self.columns.values()):
            matching = (record for record in properties.values()
                        if getattr(record, 'excel_param_name', None) == row_key)
            property_records += self._queue_deleted_records(matching, seen)

        logger.info(
            "%s: queued row deletion row_key=%r cell_records=%d property_records=%d "
            "save_items_before=%d save_items_after=%d",
            type(self).__name__, row_key, cell_records, property_records, before,
            len(self.need_update))
        return cell_records, property_records

    def queue_column_deletion(self, column_key):
        """Queue all existing sparse cell and metadata records for a column."""
        before = len(self.need_update)
        seen = set()
        cell_records = 0
        property_records = 0

        for (record_row, record_column), properties in list(self.table.items()):
            matching = (properties.values() if record_column == column_key else
                        (record for record in properties.values()
                         if getattr(record, 'date_time_izm', None) == column_key))
            marked = self._queue_deleted_records(matching, seen)
            if record_row is None or record_row not in self.ord_rows:
                property_records += marked
            else:
                cell_records += marked

        for (_, record_column), properties in list(self.columns.items()):
            matching = (properties.values() if record_column == column_key else
                        (record for record in properties.values()
                         if getattr(record, 'date_time_izm', None) == column_key))
            property_records += self._queue_deleted_records(matching, seen)

        # Include unusual metadata records without assuming where a plugin stores them.
        for properties in list(self.rows.values()):
            matching = (record for record in properties.values()
                        if getattr(record, 'date_time_izm', None) == column_key)
            property_records += self._queue_deleted_records(matching, seen)

        logger.info(
            "%s: queued column deletion column_key=%r cell_records=%d property_records=%d "
            "save_items_before=%d save_items_after=%d",
            type(self).__name__, column_key, cell_records, property_records, before,
            len(self.need_update))
        return cell_records, property_records

    def removeRow(self, row):
        if not 0 <= row < len(self.ord_rows):
            return False
        row_name = self.ord_rows[row]
        self.model().beginResetModel()
        self.ord_rows.pop(row)
        for key in [key for key in self.rows if key[0] == row_name]:
            self.rows.pop(key, None)
        for key in [key for key in self.table if key[0] == row_name]:
            self.table.pop(key, None)
        self.model()._items = {}
        self.model()._rebuild_indexes()
        self.model().endResetModel()
        return True

    def removeColumn(self, column):
        if not 0 <= column < len(self.ord_columns):
            return False
        column_name = self.ord_columns[column]
        self.model().beginResetModel()
        self.ord_columns.pop(column)
        for key in [key for key in self.columns if key[1] == column_name]:
            self.columns.pop(key, None)
        for key in [key for key in self.table if key[1] == column_name]:
            self.table.pop(key, None)
        self.model()._items = {}
        self.model()._rebuild_indexes()
        self.model().endResetModel()
        return True

    def hide_secret_rows(self):
        for row in range(self.model().rowCount()):
            self.setRowHidden(row, self.model().is_secret_row(row))

    def apply_filters(self, filters):
        if not filters:
            return
        if isinstance(filters, str):
            try:
                filters = ast.literal_eval(filters)
            except (ValueError, SyntaxError):
                logger.warning('Не удалось разобрать сохраненные фильтры: %s', filters)
                return
        if not isinstance(filters, Iterable):
            logger.warning('Некорректный формат фильтров: %s', filters)
            return

        missing_rows = set()
        for filter_data in filters:
            if not isinstance(filter_data, dict):
                continue
            x_name, y_name = filter_data.get('x'), filter_data.get('y')
            x_row, y_row = self.model()._row_indexes.get(x_name), self.model()._row_indexes.get(y_name)
            if x_row is None or y_row is None:
                missing_rows.update(str(name) for name, row in ((x_name, x_row), (y_name, y_row))
                                    if name is not None and row is None)
                continue
            for column in range(self.model().columnCount()):
                broken = self.is_value_broken(column, x_row, y_row, filter_data.get('condition'),
                                              filter_data.get('condition_percent'))
                self.item(x_row, column, create=True).update_cell('broken', str(broken))
        if missing_rows:
            logger.warning('Не удалось применить фильтры: отсутствуют строки %s', ', '.join(sorted(missing_rows)))

    def is_value_broken(self, column, x_row, y_row, condition, condition_percent):
        try:
            x_item = self.item(x_row, column)
            y_item = self.item(y_row, column)
            if x_item is None or y_item is None:
                return False
            x_value = float(x_item.value())
            y_value = float(y_item.value())
            percent = float(condition_percent)
        except (TypeError, ValueError):
            return False
        if condition == '>':
            return x_value > y_value * (1 + percent / 100)
        if condition == '<':
            return x_value < y_value * (1 - percent / 100)
        if condition == '=':
            return x_value == percent
        return False

    def clear_highlights(self):
        for item in self._highlighted_items:
            item.set_highlighted(False)
        self._highlighted_items = []

    def highlight_formula_references(self, item):
        self.clear_highlights()
        if item is None:
            return
        item.init_cells_in_formula()
        for row_name, column_name in item.cells_in_formula:
            row = self.model()._row_indexes.get(row_name)
            column = self.model()._column_indexes.get(column_name)
            if row is None or column is None:
                continue
            referenced_item = self.item(row, column)
            referenced_item.set_highlighted(True)
            self._highlighted_items.append(referenced_item)

    def format_cell_reference(self, item):
        if item is None:
            return ''
        column = self.model()._column_indexes.get(item.key[1])
        return '' if column is None else f'"{item.key[0]}"[{column + 1}]'

    def filter_table(self):
        search = (self.search_string or '').lower()
        for row, name in enumerate(self.ord_rows):
            self.setRowHidden(row, search not in name.lower())

    def copy_to_clipboard(self):
        indexes = self.selectedIndexes()
        if not indexes:
            return
        by_row = {}
        for index in indexes:
            by_row.setdefault(index.row(), []).append(index.column())
        lines = []
        for row in sorted(by_row):
            values = [self.ord_rows[row]]
            values.extend((item.text() if (item := self.item(row, column)) else '')
                          for column in sorted(by_row[row]))
            lines.append('\t'.join(values))
        QApplication.clipboard().setText('\n'.join(lines))

    def keyPressEvent(self, event):
        if event.key() == Qt.Key_F:
            search, ok = QInputDialog.getText(self, 'Поиск', 'Введите строку для поиска:')
            if ok:
                self.search_string = search
                self.filter_table()
            return
        if event.matches(QKeySequence.Copy):
            self.copy_to_clipboard()
            return
        super().keyPressEvent(event)

    def mousePressEvent(self, event):
        parent = self._parent
        if parent and hasattr(parent, 'formula_edit') and parent.formula_edit.hasFocus():
            index = self.indexAt(event.pos())
            if index.isValid():
                reference = self.format_cell_reference(self.itemFromIndex(index))
                if reference:
                    parent.begin_formula_reference_capture()
                    self.setCurrentIndex(index)
                    parent.formula_edit.insert_reference(reference)
                    QTimer.singleShot(0, parent.formula_edit.setFocus)
                    QTimer.singleShot(0, parent.end_formula_reference_capture)
                    return
        super().mousePressEvent(event)
